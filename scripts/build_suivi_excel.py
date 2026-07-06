#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère le classeur de suivi d'avancement du carnet étanchéité v02,
sur le même principe que le tableau récapitulatif de détails fourni en modèle
(colonnes N°/Intitulé/Modification/Avancement/Commentaires, statuts colorés).

Source de vérité du contenu : docs/suivi_reprises.md + NOTES.md + historique git.

Usage: python3 build_suivi_excel.py OUT.xlsx
"""
import sys
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]

# ------------------------------------------------------------------ styles
FONT_NAME = "Calibri"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_TOP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")

STATUS_STYLES = {
    "Fait":                ("C6EFCE", "006100"),
    "Fait partiellement":  ("FFEB9C", "9C6500"),
    "Déjà conforme":       ("C6EFCE", "006100"),
    "À faire":             ("FFC7CE", "9C0006"),
    "À faire (création)":  ("FFC7CE", "9C0006"),
    "À faire (dessin)":    ("D9D9D9", "3B3B3B"),
    "Bloqué":              ("C00000", "FFFFFF"),
    "Question":            ("DDEBF7", "1F4E78"),
    "Dessin":              ("D9D9D9", "3B3B3B"),
    "Référence":           ("F2F2F2", "595959"),
}

def style_status(cell, status):
    bg, fg = STATUS_STYLES.get(status, ("FFFFFF", "000000"))
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=fg)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.value = status

def write_header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = WRAP_TOP_CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate

def write_title(ws, text, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = SUBTITLE_FONT
    ws.row_dimensions[2].height = 16

def fill_row(ws, row, values, wraps=None):
    wraps = wraps or {}
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.border = BORDER
        c.alignment = wraps.get(i, WRAP_TOP)
        c.font = Font(name=FONT_NAME, size=10)
    return row


# ============================================================ Feuille 1
GEN_DATE = datetime.date(2026, 7, 6)
BRANCH = "claude/carnet-etancheite-maj-f67-sshiyr"
COMMITS = "c503e4c -> fb67755 (5 commits, 05-06/07/2026)"

wb = openpyxl.Workbook()
ws1 = wb.active
ws1.title = "Suivi général"
ncols = 8
write_title(
    ws1,
    "Carnet de details étanchéité OS — Suivi des reprises v02 (MAJ F67-III / AFTES GT9)",
    f"Généré le {GEN_DATE.strftime('%d/%m/%Y')} — branche {BRANCH} — commits {COMMITS} — "
    f"Référence du mail : docs/mail_bertrand_reprises_v02.txt",
    ncols,
)
HEADERS1 = ["Réf.", "Sujet / Folio", "Modification à apporter (mail Bertrand)",
            "Modification apportée", "Date de reprise", "Avancement",
            "Référence technique", "Commentaire"]
WIDTHS1 = [10, 30, 48, 48, 14, 16, 22, 40]
write_header(ws1, 4, HEADERS1, WIDTHS1)

D1, D2 = "05/07/2026", "06/07/2026"

# (ref, sujet, mod_a_apporter, mod_apportee, date, statut, reference, commentaire)
ROWS = [
 ("T1", "Transverse — Folio 00 references",
  "Ajouter un folio 00 «Références et généralités» : F67-III, GT9.C1F1, GT9.C2, "
  "GT9.R19F1, NF EN 13491/13256, AT CETU, ASQUAL soudeurs, réception support = point d'arrêt.",
  "", "", "À faire (création)", "-", "Valeurs de contrôle C1 p.81 déjà rassemblées dans le mail"),
 ("T2", "Transverse — Légende complexe DEG",
  "Harmoniser sur tous les folios : géotextile inférieur / géomembrane PVC-P translucide "
  "2mm mini / écran PVC supérieur.",
  "Harmonisation partielle possible ; épaisseur écran supérieur non tranchée.",
  "", "Fait partiellement", "-", "Épaisseur écran sup. = point bloquant n°2"),
 ("T3", "Transverse — Mortiers R3/R4",
  "Préciser partout «classe R3/R4 selon NF EN 1504-3».",
  "27 occurrences «NF EN 1504-3» dans le fichier (mortiers, ragréages, resurfaçages).",
  D1, "Fait", "Lot 1 (2af7223)", ""),
 ("T4", "Transverse — Joint hydrogonflant",
  "Harmoniser hydroexpansif / hydrogonflant / hydro-expansif -> «joint hydrogonflant».",
  "41 occurrences harmonisées, 0 résidu hydro(-)expansif (vérifié par grep brut).",
  D1, "Fait", "Lot 1 (2af7223)", ""),
 ("T5", "Transverse — Enduit protection au feu",
  "Ajouter renvoi référentiel exploitant + PV feu produit.",
  "Renvoi ajouté en ligne sur folios 9/10 (+2 multileaders folio 65). Légendes empilées "
  "des folios 11-19/24 non touchées (risque de chevauchement graphique).",
  D2, "Fait partiellement", "Lot 4 (fb67755)", "Prescription générale à porter au futur folio 00"),

 ("P1", "Folio 1 — Vue puits PM / rameau DEG",
  "Harmoniser la casse des renvois «Détail n°1/2/3» avec les folios 5, 26, 28.",
  "Renvois alignés sur le folio 5 : Partie supérieure / Piédroit / Radier.",
  D2, "Fait", "Lot 4, #FDF7", ""),
 ("P2", "Folio 2 — Puits PM / rameau DEG, partie sup.",
  "Mortier R4 -> NF EN 1504-3. Retour technique 1,00 m : ajouter «à adapter selon "
  "hauteur d'eau / note de calcul».",
  "Les deux reprises appliquées.",
  D1, "Fait", "Lot 1 (generique) + Lot 2, #1FE5F", ""),
 ("P4/P20", "Folios 4 et 20 — Solin 3x3 cm",
  "Harmoniser «solin 3x3cm mortier époxy» (F4) avec «mortier polymérique» (F20).",
  "", "", "Question", "-", "Le mail ne tranche pas la cible : époxy ou polymérique ?"),
 ("P8", "Folio 8 — Radier SELA",
  "Renvoyer explicitement le «Étanchéité DEG (voir détail...)» vers les folios 30-33.",
  "Renvoi «(voir détails folios 30 à 33)» ajouté.",
  D1, "Fait", "Lot 2, #12A34", ""),
 ("P9", "Folio 9 — Jonction tunnel trad. SELA / tunnel foré",
  "Mortiers R4/retrait compensé -> NF EN 1504-3. Ajouter note préparation support "
  "voussoir (surfaçage épaufrures, traitement joints, C1 §6.4).",
  "Norme ajoutée + NOTA de préparation support voussoir créé.",
  D1, "Fait", "Lot 1 + Lot 3, #20AC4", ""),
 ("P10", "Folio 10 — Jonction tunnel trad. DEG / tunnel foré",
  "Prévoir un folio dédié pour l'arrêt du DEG sur voussoir (bride/contre-bride ou tole "
  "colaminée), traité aujourd'hui par un simple solin.",
  "", "", "À faire (création)", "-", "Folio à créer n°2 de la liste"),
 ("P11,16-19", "Folios 11, 16-19 — Vues de principe",
  "Libellés des détails en corps trop petit et casse hétérogène : harmoniser.",
  "", "", "Question", "-", "Pas d'anomalie objective au rendu ; critère de cible à préciser"),
 ("P12-14 a", "Folios 12-14 — Dalle intermediaire / piedroit",
  "Légender le complexe «Étanchéité par D.E.G.».",
  "Légende déjà présente dans la source sur les folios 12/13/14/15 (source plus fraîche "
  "que le PDF v01 retouché).",
  "-", "Déjà conforme", "-", ""),
 ("P12-14 b", "Folios 12-14 — Cote 0,50 m",
  "Cote 0,50 m côté support : préciser sa signification (largeur de préparation ?).",
  "", "", "Question", "-", "Ambiguïté signalée par le mail lui-même"),
 ("P15", "Folio 15 — Radier",
  "Vérifier présence de l'écran de protection supérieure dans la légende côté gauche "
  "(non légendé).",
  "MTEXT «Écran de protection supérieure» ajouté en légende gauche.",
  D2, "Fait", "Lot 4, #20FA3", ""),
 ("P20", "Folio 20 — Radier",
  "«Solin 3x3 mortier polymérique» : harmoniser avec folio 4. Surfaçage époxy/R4 -> "
  "NF EN 1504-3.",
  "Norme NF EN 1504-3 appliquée. Harmonisation solin : voir question P4/P20.",
  D1, "Fait", "Lot 1", ""),
 ("P21", "Folio 21 — Vue tunnel fore / gare SELA",
  "Vérifier cohérence «Min 1.00m» avec les cotes des folios 22-25.",
  "«retour de 1m» harmonisé en «retour de 1.00m» (3 multileaders).",
  D2, "Fait partiellement", "Lot 4", "Vérification des cotes 22-25 vs principe : revue Alexandre"),
 ("P22/23", "Folios 22-23 — Dalle intermediaire",
  "Mortier ragréage R4 -> NF EN 1504-3. «Découpe possible du voussoir» : ajouter note "
  "validation structure (accord MOE GC obligatoire).",
  "Note «(accord préalable MOE GC obligatoire)» ajoutée (3 occurrences).",
  D1, "Fait", "Lot 1", ""),
 ("P24", "Folio 24 — Piedroit",
  "Doublon de légende «Joint hydroexpansif» (x2) : nettoyer.",
  "Doublon supprimé (#EB7E, superposé à 0.00 unité près de #BA4F).",
  D1, "Fait", "Lot 1", ""),
 ("P25", "Folio 25 — Radier",
  "NOTA «Traitement des joints de PM à prévoir en radier» : renvoyer explicitement aux "
  "folios 40-43.",
  "Renvoi «(cf. folios 40 à 43)» ajouté.",
  D1, "Fait", "Lot 2, #6651", ""),
 ("P27 a", "Folio 27 — Jonction tunnel fore / rameau SELA",
  "Seuil «si hauteur d'eau > 10 m» : sourcer, sinon écrire «selon note de calcul».",
  "", "", "Bloqué", "-", "Point bloquant n°5 — règle interne ou à sourcer"),
 ("P27 b", "Folio 27 — idem",
  "Harmoniser tole galva/inox avec folios 29 et 65.",
  "«tole (galvanisée ou inox)» + coquille «conformtement» corrigée (seuil >10m intact).",
  D1, "Fait", "Lot 2, #FEBB", ""),
 ("P29", "Folio 29 — Jonction tunnel fore / rameau DEG",
  "Plats galvanisés ép.4mm : option inox possible. Système bride/contre-bride : "
  "entraxe fixations, couple serrage NDC, nature joints compressibles.",
  "«galvanisés ou inox» + NOTA bride/contre-bride créé.",
  D1, "Fait", "Lot 1, #1637D + Lot 3, #20B9B", ""),
 ("P30/32", "Folios 30, 32 — DEG sous radier ancre PM",
  "Géotextile de protection à préciser (renvoi GT9.R19F1). Cunette : rattacher au "
  "drainage GT9.R19F1.",
  "Cunettes -> «(drainage selon GT9.R19F1)» (2 multileaders). Précision géotextile liée "
  "au futur folio 00 / T2.",
  D2, "Fait partiellement", "Lot 4", ""),
 ("P31/33", "Folios 31, 33 — DEG sous radier, ancrage",
  "Ajouter continuité de l'enduit d'imperméabilisation jusqu'au trait de scie + phasage.",
  "NOTA créé sur les deux folios.",
  D1, "Fait", "Lot 3, #20AC6/#20AC7", ""),
 ("P34 a", "Folio 34 — Interface DEG / pieu prefonde",
  "Ajouter le phasage en NOTA (recépage, coffrage, béton hydrofuge, pose profilés, "
  "injection).",
  "NOTA phasage créé.",
  D1, "Fait", "Lot 3, #20B99", ""),
 ("P34/36 b", "Folios 34, 36 — Chape de protection DEG",
  "Chape de protection DEG : 5 cm ici vs 6 cm au folio 36. Harmoniser.",
  "Chape #6030 (F34) 5->6 cm ; #6071 déjà à 6 cm. Vérifié au lot 4 : plus aucune chape "
  "DEG à 5 cm.",
  f"{D1} / {D2}", "Fait", "Lot 1, #6030 (verif. Lot 4)", ""),
 ("P37", "Folio 37 — Traversée bride et contre-bride",
  "Comme folio 29 : entraxes de fixation, joints compressibles, résine de protection à "
  "compléter.",
  "NOTA créé (même contenu que P29).",
  D1, "Fait", "Lot 3, #20B9A", ""),
 ("P38", "Folio 38 — Traitement PM, dessaffleurement",
  "Cotes engravures 9/6/10/15 cm : à vérifier contre le CMO du procédé retenu.",
  "", "", "Bloqué", "-", "Point bloquant n°3"),
 ("P39", "Folio 39 — Arrêt étanchéité tole colaminée",
  "Soudure manuelle sur tole colaminée : ajouter contrôle à la pointe sèche ou cloche à "
  "vide (C1 p.81).",
  "NOTA contrôle créé.",
  D1, "Fait", "Lot 3, #20AC9", ""),
 ("P41", "Folio 41 — Joint PM, traitement confortatif",
  "Bande de pontage élastomérique 200mm collée résine : préciser résine (époxy) + "
  "préparation.",
  "«collée à la résine époxy (préparation du support selon GT9.C1F1)» (4 occurrences, "
  "même objet folios 41/43/65).",
  D1, "Fait", "Lot 2", ""),
 ("P44", "Folio 44 — DEG sur paroi berlinoise",
  "Intégrer les exigences de préparation berlinoise (désaffleurements <=5cm, chanfrein "
  "45°, remplissage ondes 90 kPa).",
  "NOTA créé avec les 3 exigences.",
  D1, "Fait", "Lot 3, #20ACA", ""),
 ("P45 a", "Folio 45 — Raccord DEG / FPM berlinoise",
  "Détail raccord DEG/FPM «à définir» depuis 2022 : à trancher.",
  "", "", "Bloqué", "-", "Point bloquant n°4"),
 ("P45 b", "Folio 45 — idem",
  "FPM bicouche anti-racine : uniquement si végétalisation (cohérence folio 61 qui "
  "montre monocouche).",
  "«(uniquement si végétalisation)» ajouté sur #1746C.",
  D1, "Fait", "Lot 2", "Même texte au folio 63 (#13E13) non modifié : voir question"),
 ("P47", "Folio 47 — FPM remontées édicules",
  "«Béton de protection ép.6 cm» : conforme, harmoniser avec les folios où 5 cm est "
  "indiqué.",
  "Sert de référence cible (6 cm) pour l'harmonisation P34/36.",
  "-", "Question", "-", "Généraliser 5->6cm partout ? voir question dédiée"),
 ("P48", "Folio 48 — FPM remontées sous niveau TN",
  "Harmoniser casse / unites.",
  "", "", "Question", "-", "Pas d'anomalie objective au rendu ; régression Lot2 corrigée au Lot4"),
 ("P49", "Folio 49 — FPM joint de dilatation",
  "Ajouter note sur le soufflet au droit du JD et la non-adhérence locale de la feuille.",
  "NOTA créé.",
  D1, "Fait", "Lot 3, #20B9C", ""),
 ("P50", "Folio 50 — FPM evacuation et drainage",
  "Protection anti-intrusion + entretenabilité à formaliser sur le dessin.",
  "", "", "À faire (dessin)", "-", "Nécessite entités graphiques, lot dédié"),
 ("P51", "Folio 51 — FPM raccord arase basse",
  "«Resurfaçage de la paroi moulée» -> mortier NF EN 1504-3 R3 mini + tolérances.",
  "Norme + tolérances GT9.C1F1 ajoutées.",
  D1, "Fait", "Lot 2, #130D3", ""),
 ("P52-54 a", "Folios 52-54 — FPM raccord arase haute",
  "«Resurfaçage soigné» -> NF EN 1504-3 R3 mini.",
  "Norme appliquée (3 occurrences).",
  D1, "Fait", "Lot 2", ""),
 ("P52-54 b", "Folios 52-54 — idem",
  "Ajouter en tête de section un tableau de choix selon position nappe/arase (C2 §4.5, "
  "niveaux EB/EH/EE).",
  "", "", "À faire (création)", "-", "Cite le Cahier 2 -> dépend du point bloquant n°1"),
 ("P55/57/59", "Folios 55, 57, 59 — FPM arase super haute",
  "NOTA carottages : renvoyer au folio 50. Resurfaçage type R3 -> NF EN 1504-3.",
  "Renvoi «(cf. folio 50)» + norme R3 ajoutés.",
  D1, "Fait", "Lot 1 + Lot 2", ""),
 ("P56-60", "Folios 56-60 — FPM arase super haute",
  "Feuillard métallique : préciser matériau (inox/galva) + entraxe (FPM verticale : "
  "4 fixations/ml tous les 3 m). Resurfaçage soigné -> NF EN 1504-3.",
  "Matériau + entraxe précisés (9 occurrences) ; norme resurfaçage appliquée (6 occ.).",
  D1, "Fait", "Lot 1 + Lot 2", ""),
 ("P61 a", "Folio 61 — Raccordement DEG/FPM avec SELA",
  "Préciser «SEL-A sous référentiel CETU».",
  "Légende et titre corrigés.",
  D1, "Fait", "Lot 2, #13DD7/#15685", ""),
 ("P61 b", "Folio 61 — idem",
  "Recouvrements 20/20/15 cm : vérifier contre les CMO des procédés.",
  "", "", "Question", "-", "Vérification externe aux CMO des procédés prescrits"),
 ("P65 a", "Folio 65 — Bandes de pontage",
  "«Engravure 20x220 mm» : cote probablement erronée, à vérifier.",
  "", "", "Question", "-", "Valeur de remplacement à fournir"),
 ("P65 b", "Folio 65 — idem",
  "«Centré 70 mm mini» : clarifier (débord mini de part et d'autre du joint).",
  "«Débord 70mm mini de part et d'autre du joint» (2 occurrences).",
  D1, "Fait", "Lot 2", ""),
 ("P65 c", "Folio 65 — idem",
  "Tole galva ou inox : harmoniser avec folios 27 et 29.",
  "Déjà conforme («galvanisé ou Inox»).",
  "-", "Déjà conforme", "-", ""),
]

r = 5
for row in ROWS:
    ref, sujet, apporter, apportee, date, statut, refe, comment = row
    fill_row(ws1, r, [ref, sujet, apporter, apportee, date],
             wraps={1: WRAP_TOP_CENTER, 5: WRAP_TOP_CENTER})
    style_status(ws1.cell(row=r, column=6), statut)
    fill_row(ws1, r, [None, None, None, None, None, None, refe, comment])
    ws1.cell(row=r, column=1).value = ref
    ws1.cell(row=r, column=7).font = Font(name=FONT_NAME, size=9, italic=True, color="595959")
    ws1.cell(row=r, column=7).alignment = WRAP_TOP
    r += 1

last_row = r - 1
for row in ws1.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=8):
    for cell in row:
        cell.border = BORDER

# ligne de synthese
r += 1
from collections import Counter
counts = Counter(row[5] for row in ROWS)
ws1.cell(row=r, column=1, value="Synthèse :").font = Font(name=FONT_NAME, bold=True)
r += 1
for label in ["Fait", "Fait partiellement", "Déjà conforme", "À faire (création)",
              "À faire (dessin)", "Bloqué", "Question"]:
    n = counts.get(label, 0)
    c1 = ws1.cell(row=r, column=1, value=label)
    c1.font = Font(name=FONT_NAME, size=10)
    c2 = ws1.cell(row=r, column=2, value=n)
    c2.font = Font(name=FONT_NAME, size=10, bold=True)
    r += 1
total = sum(counts.values())
c1 = ws1.cell(row=r, column=1, value="Total items suivis")
c1.font = Font(name=FONT_NAME, size=10, bold=True)
ws1.cell(row=r, column=2, value=total).font = Font(name=FONT_NAME, size=10, bold=True)

ws1.sheet_view.zoomScale = 100


# ============================================================ Feuille 2 : Points bloquants
ws2 = wb.create_sheet("Points bloquants")
ncols2 = 5
write_title(ws2, "Points bloquants — décision Bertrand Verrière requise",
            "Aucun de ces points n'a été tranché automatiquement (règle CLAUDE.md du projet).",
            ncols2)
HEADERS2 = ["N°", "Sujet", "Description", "Folio(s) concerné(s)", "Décision requise de"]
WIDTHS2 = [6, 30, 70, 18, 22]
write_header(ws2, 4, HEADERS2, WIDTHS2)

BLOQUES = [
 (1, "Statut du Cahier 2 GT9",
  "Cahier 2 (guide de choix, prépublication septembre 2025) : statut non confirmé. "
  "Ne pas citer comme référence contractuelle sans validation.",
  "Folio 00 (à créer), P52-54 (tableau de choix)", "Bertrand Verrière"),
 (2, "Épaisseur écran supérieur DEG",
  "Le GT9 C1F1 exige 19/10e mini, le carnet v01 indique 20/10e à plusieurs endroits. "
  "Confirmer la valeur cible pour la v02.",
  "Légende complexe DEG (tous folios DEG)", "Bertrand Verrière"),
 (3, "Cotes d'engravure du folio 38",
  "Valeurs actuelles (9/6/10/15 cm) à comparer aux CMO des procédés habituellement "
  "prescrits.",
  "Folio 38", "Bertrand Verrière"),
 (4, "Raccord DEG/FPM du folio 45",
  "Marqué «à définir» depuis la v01 (2022) : à trancher.",
  "Folio 45", "Bertrand Verrière"),
 (5, "Seuil «hauteur d'eau > 10 m» du folio 27",
  "Règle interne à documenter, ou valeur à sourcer dans un référentiel.",
  "Folio 27", "Bertrand Verrière"),
]
r = 5
for n, sujet, desc, folios, qui in BLOQUES:
    fill_row(ws2, r, [n, sujet, desc, folios, qui],
             wraps={1: WRAP_TOP_CENTER, 4: WRAP_TOP_CENTER, 5: WRAP_TOP_CENTER})
    for c in range(1, ncols2 + 1):
        ws2.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FCE4E4")
    r += 1
ws2.sheet_view.zoomScale = 100


# ============================================================ Feuille 3 : Questions Alexandre
ws3 = wb.create_sheet("Questions Alexandre")
ncols3 = 4
write_title(ws3, "Questions ouvertes pour Alexandre",
            "Ambiguës mais non bloquantes au sens CLAUDE.md — réponse d'Alexandre suffisante "
            "(pas nécessairement Bertrand).",
            ncols3)
HEADERS3 = ["N°", "Sujet", "Question", "Folio(s)"]
WIDTHS3 = [6, 26, 70, 18]
write_header(ws3, 4, HEADERS3, WIDTHS3)

QUESTIONS = [
 (1, "Solin 3x3 cm", "Mortier époxy (folio 4) ou polymérique (folio 20) : quelle cible unique ?", "P4, P20"),
 (2, "Béton de protection lourde 5 cm",
  "Le folio 47 valide 6 cm pour les remontées d'édicules. Faut-il généraliser 5->6 cm "
  "partout (~12 occurrences, folios 15/33/34/40-43/46/48/51...) ou est-ce un autre "
  "élément (protection lourde horizontale) qui reste à 5 cm ?", "P15, P33, P34, P40-43, P46, P48, P51..."),
 (3, "Cote 0,50 m", "Signification à confirmer avant de la légender.", "P12-14"),
 (4, "Engravure folio 65", "«20x220 mm» signalée comme probablement erronée : valeur de remplacement ?", "P65"),
 (5, "Anti-racine folio 63",
  "Même texte qu'au folio 45 (bicouche anti-racine) : appliquer la même condition "
  "«uniquement si végétalisation» ? Et le titre du folio 61 indique «D.E.G. - F.P.B» "
  "(F.P.B au lieu de F.P.M ailleurs) : coquille à confirmer.", "P63, P61"),
 (6, "Convention décimale",
  "La source utilise le point partout (197 occurrences). La correction PDF «1,00m» "
  "impliquerait de basculer ~41 cotes en virgule. Bascule générale à trancher (charte "
  "Egis) avant toute modification massive.", "Transverse"),
 (7, "Casse/corps des libellés",
  "Folios 11, 16-19 : casse hétérogène signalée par le mail, mais pas d'anomalie "
  "objective identifiée au rendu. Quel est le critère de cible ?", "P11, P16-19"),
 (8, "Casse/unites folio 48",
  "Idem : pas d'anomalie objective au rendu (une régression du lot 2 sur le feuillard a "
  "été corrigée au lot 4). Cible d'harmonisation à préciser.", "P48"),
]
r = 5
for n, sujet, question, folios in QUESTIONS:
    fill_row(ws3, r, [n, sujet, question, folios],
             wraps={1: WRAP_TOP_CENTER, 4: WRAP_TOP_CENTER})
    for c in range(1, ncols3 + 1):
        ws3.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FFF9E0")
    r += 1
ws3.sheet_view.zoomScale = 100


# ============================================================ Feuille 4 : Folios a creer
ws4 = wb.create_sheet("Folios à créer")
ncols4 = 4
write_title(ws4, "Nouveaux folios à créer (source uniquement)",
            "Non réalisable en retouche PDF ; nécessite la source DXF/DWG.",
            ncols4)
HEADERS4 = ["N°", "Folio", "Description", "Avancement"]
WIDTHS4 = [6, 32, 60, 18]
write_header(ws4, 4, HEADERS4, WIDTHS4)

NOUVEAUX = [
 (1, "Folio 00 — Références et généralités",
  "F67-III, GT9.C1F1, GT9.C2 (sous réserve), GT9.R19F1, NF EN 13491/13256, AT CETU, "
  "ASQUAL soudeurs, réception support = point d'arrêt, valeurs de contrôle C1 p.81.",
  "A faire"),
 (2, "Arrêt d'étanchéité DEG sur voussoir",
  "Bride/contre-bride ou tole colaminée — aujourd'hui traité par un simple solin "
  "(folio 10).", "A faire"),
 (3, "Détails drainage",
  "Géocomposite piédroit, barbacane, cunette raccordée, selon GT9.R19F1.", "A faire"),
 (4, "Traversées de réseaux courants",
  "Fourreaux, manchettes — hors bride/contre-bride.", "A faire"),
 (5, "Tableau de choix DEG / SEL-A / FPM",
  "Par partie d'ouvrage et hauteur d'eau, inspiré du Cahier 2 §4.5 et §8.",
  "Bloqué (statut Cahier 2)"),
 (6, "Réparation / réinjection d'un compartiment DEG",
  "Détail non existant dans le carnet actuel.", "A faire"),
]
r = 5
for n, folio, desc, av in NOUVEAUX:
    fill_row(ws4, r, [n, folio, desc, None], wraps={1: WRAP_TOP_CENTER})
    style_status(ws4.cell(row=r, column=4), "Bloqué" if "Bloqué" in av else "À faire (création)")
    r += 1
ws4.sheet_view.zoomScale = 100

wb.save(OUT)
print(f"Classeur genere : {OUT}")
print(f"Feuilles : {wb.sheetnames}")
print(f"Lignes suivi general : {len(ROWS)}")
