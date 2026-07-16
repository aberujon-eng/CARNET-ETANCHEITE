import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
p = "docs/Carnet_etancheite_suivi_avancement_v02.xlsx"
wb = openpyxl.load_workbook(p)
name = "MAJ normes 07-2026"
if name in wb.sheetnames:
    del wb[name]
ws = wb.create_sheet(name, 1)
hdr = Font(bold=True, color="FFFFFF", size=11)
fill = PatternFill("solid", fgColor="2F5496")
ok = PatternFill("solid", fgColor="C6EFCE")
res = PatternFill("solid", fgColor="FFF2CC")
thin = Border(*[Side(style="thin", color="999999")]*4)
wrap = Alignment(wrap_text=True, vertical="top")
ws.append(["Folio(s)", "Modification", "Justification / Norme", "Statut"])
for c in ws[1]:
    c.font = hdr; c.fill = fill; c.border = thin
rows = [
    ("15/33/34/40-43/46/48/51/61-63", "Béton de protection lourde 5 cm -> 6 cm (34 occ.)", "GT9.C5F1 (prépub.) tableaux 1-3 ; cohérent P34/36 + P47 mail", "FAIT"),
    ("30", "NOTA pipettes : >=3/compartiment en boîtier poteau/refend", "GT9.C10F1 §7.1 (publication)", "FAIT"),
    ("32", "NOTA surfaces maxi compartimentage 250/350 m²", "F67-III art. 3.3.1", "FAIT"),
    ("46", "NOTA débit AFTES niveau 0 + retombée >=20 cm", "GT9.C10F1 §9", "FAIT"),
    ("50", "Grille anti-intrusion démontable + entretenabilité (graphique)", "P50 mail ; GT9.C10F1 §9.3.1", "FAIT"),
    ("00 (créé)", "Folio Références et généralités + contrôles C1 p.81", "Mail Bertrand ; GT9.C1F1 p.81 ; F67-III §3.1-3.2", "FAIT"),
    ("66 (créé)", "Arrêt DEG sur voussoir (tôle colaminée / bride-contre-bride)", "GT9.C9F1 §6.5 (prépub.) ; F67-III §3.1.3.c", "FAIT"),
    ("67 (créé)", "Détails drainage (nappe, cunette, collecteur)", "GT9.R19F1 ; GT9.C10F1 §10.2", "FAIT"),
    ("68 (créé)", "Traversées de réseaux courants (fourreau + hydrogonflant)", "GT9.C10F1 §8 ; F67-III §3.1.7", "FAIT"),
    ("69 (créé)", "Tableau de choix DEG/SEL-A/FPM — EN RÉSERVE", "Dépend statut Cahier 2 = point BLOQUÉ n°1 (B. Verrière)", "RÉSERVE"),
    ("70 (créé)", "Réparation/réinjection compartiment DEG (procédure)", "GT9.R1F2 III.1.5.2.4", "FAIT"),
]
for r in rows:
    ws.append(r)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.border = thin; c.alignment = wrap
        if c.column == 4:
            c.fill = ok if c.value == "FAIT" else res
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 52
ws.column_dimensions["C"].width = 50
ws.column_dimensions["D"].width = 10
wb.save(p)
print("feuille ajoutée:", name, "-", len(rows), "lignes")
